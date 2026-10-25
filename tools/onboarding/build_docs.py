from pathlib import Path
import json, shutil, math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RImage, PageBreak, Table, TableStyle
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
root=Path(__file__).resolve().parents[2];out=root/'output/pdf/onboarding-0.13';src=out/'screenshots/source';ann=out/'screenshots/annotated';ann.mkdir(exist_ok=True)
base=root/'examples/manual-baseline'
BLUE='225DA8';TEAL='427E70';AMBER='A46D24';INK='252A2E'
font='/System/Library/Fonts/Supplemental/Arial.ttf'
fontb='/System/Library/Fonts/Supplemental/Arial Bold.ttf'
def asset(name, file, crop=None, marks=(), baseline=False):
 im=Image.open((base if baseline else src)/file).convert('RGB')
 if crop:im=im.crop(crop)
 d=ImageDraw.Draw(im); f=ImageFont.truetype(fontb,26)
 for num,x,y,tx,ty in marks:
  d.line((x,y,tx,ty),fill='#4B8AD5',width=4)
  a=math.atan2(ty-y,tx-x)
  d.polygon([(tx,ty),(tx-15*math.cos(a-.5),ty-15*math.sin(a-.5)),(tx-15*math.cos(a+.5),ty-15*math.sin(a+.5))],fill='#4B8AD5')
  d.ellipse((x-22,y-22,x+22,y+22),fill='#225DA8',outline='white',width=2)
  d.text((x,y),str(num),font=f,fill='white',anchor='mm')
 p=ann/(name+'.png');im.save(p);return str(p)
A={}
A['workspace']=asset('workspace','workspace.png',marks=[(1,160,1050,60,1100),(2,1400,600,1600,980),(3,2800,280,3150,280),(4,2780,900,3250,900)])
# Actual native screenshot crops preserve UI pixels and aspect ratio.
A['modules']=asset('modules','workspace.png',(3020,690,3590,970))
A['context']=asset('context','pinned.png',(3020,690,3590,1180))
A['install']=asset('install','final-installed-enabled.png')
A['search']=asset('search','workspace-search-fixed.png',(1500,1400,2550,1910))
A['screen']=asset('screen','first-screen-controls.png',(3010,780,3590,1510))
A['projector']=asset('projector','solver-before.png',(3020,950,3590,1390))
A['calc']=asset('calculated','solver-six-metres.png',(3020,1340,3590,1840))
A['solver']=asset('solver','solver.png',(3020,1310,3590,2020))
A['training-solve']=asset('training-solve','exercise-solver.png',(3020,1380,3590,1840))
A['optics']=asset('optics','optics-inspector.png',(65,890,625,1290),baseline=True)
A['blend']=asset('blend','exercise-reflow.png',(3020,970,3590,1660))
A['auto']=asset('auto','auto-blend-valid.png',(2510,1260,3425,2135))
A['views']=asset('views','views.png',(3020,900,3590,1600))
A['density']=asset('density','exercise-density-after.png',(3020,1000,3590,1750))
A['before']=asset('before','exercise-density-before.png',(700,650,2600,1400))
A['after']=asset('after','exercise-density-after.png',(700,650,2600,1400))
A['stretch']=asset('stretch','stretch-asset.png')
A['reverse']=asset('reverse','exercise-reverse-dim.png',(700,790,1860,1160))
A['covered']=asset('covered','exercise-reverse-covered.png',(700,790,1860,1160))
A['hidden']=asset('hidden','exercise-reverse-hidden.png',(700,790,1860,1160))
A['overlap']=asset('overlap','exercise-reverse-overlap.png',(700,790,1860,1160))
A['export']=asset('export','exercise-export.png',(3020,740,3590,1250),[(1,45,60,265,75),(2,45,115,265,125),(3,45,265,265,275)])
A['export-density']=asset('export-density','exercise-density-export-settings.png',(3020,1350,3590,1860))
A['theme']=asset('theme','theme-confirm.png',(2660,1370,3590,1900))
A['install-menu']=asset('install-menu','final-install-menu.png',(1300,65,1670,240))
im=Image.open(A['install']);menu=Image.open(A['install-menu']);menu=menu.resize((740,350));im.paste(menu,(700,560));im.save(A['install'])
A['reverse-controls']=asset('reverse-controls','exercise-reverse-dim.png',(3020,1790,3590,2130))
A['density-overlay']=asset('density-overlay','../'+'../Beam_Training_Project_01_Density.png')
A['overview']=asset('overview','workspace.png',(250,500,2800,1750))
im1=Image.open(A['before']);im2=Image.open(A['after'])
compare=Image.new('RGB',(im1.width,im1.height+im2.height+140),'#F4F5F6');d=ImageDraw.Draw(compare);ff=ImageFont.truetype(fontb,34)
d.text((20,12),'BEFORE · 93.1% visible coverage',fill='#252A2E',font=ff);compare.paste(im1,(0,70))
d.text((20,im1.height+82),'AFTER · 78.6% visible coverage',fill='#252A2E',font=ff);compare.paste(im2,(0,im1.height+140))
A['comparison']=str(ann/'occlusion-comparison.png');compare.save(A['comparison'])
mode_keys=['covered','reverse','hidden','overlap'];labels=['Covered Only','Uncovered Dimmed','Uncovered Hidden','Overlap Diagnostic']
ww=1160;hh=420;grid=Image.new('RGB',(ww*2,hh*2),'#F4F5F6');gd=ImageDraw.Draw(grid);gf=ImageFont.truetype(fontb,32)
for i,(key,label) in enumerate(zip(mode_keys,labels)):
 x=(i%2)*ww;y=(i//2)*hh;gd.text((x+16,y+8),label,font=gf,fill='#252A2E');grid.paste(Image.open(A[key]),(x,y+48))
A['mode-grid']=str(ann/'reverse-modes.png');grid.save(A['mode-grid'])
# Each page is an operational topic with explicit page breaks, readable detail crops and a short checkpoint.
pages=[]
def page(title,intro,blocks=(),img=None,caption='',width=3.2):pages.append(dict(title=title,intro=intro,blocks=blocks,img=img,caption=caption,width=width))
page('Beam User Guide','Twisted Pixel • Projection planning in Blender\nBeam 0.13.0 — Internal Production Beta',[
('p','Beam helps you place projectors, size images, plan blends and inspect coverage on scene geometry. Use it to build a technical study before installation, then export the study for review or downstream tools.'),
('p','This guide starts with the workspace, then takes you from a new screen to analysis and handoff. Open the supplied training project when you reach the exercises. Basic Blender orbit, pan and object selection are assumed.'),
('p','Workspace map: 1 Module rail • 2 3D View • 3 Outliner • 4 Beam Inspector. The next pages enlarge the controls.'),
('note','Tested with Blender 5.2.2 LTS on macOS 26.7. This is an internal beta in testing, not a stable public release.')],A['workspace'],'Actual Beam workspace with the supplied three-projector training study.',6.8)
page('Find the Beam tools','Beam lives inside Blender. The Beam workspace arranges the 3D View, Outliner and a contextual inspector for projection work.',[
('table',[['Module','Use it for'],['Projector','Create PJ objects; resolution, lens, aim, calculations and solver.'],['Screen','Create SCR surfaces; dimensions, aspect and appearance.'],['Blend','Create and edit BG groups, overlap and member placement.'],['Views','Save, recall and export study framing.'],['Export','Projector data, study packages, images and screen geometry.']]),
('p','The module rail sits below Blender’s native tools at the left of the 3D View. A narrow toolbar shows icons. Drag its right edge wider to reveal names. Press T with the pointer over the 3D View if the toolbar is hidden.'),
('p','The same five module buttons are at the top of the Beam Inspector. Use those if the rail is hidden. Frame, Solo and Lock also remain available in the inspector.'),
('p','Use the Outliner at upper right to select exact objects: BEAM → Projectors, Screens and blend controllers. Expand the triangles beside collections. Selection here is often easier than clicking overlapping geometry.')],A['modules'],'The inspector repeats the modules and shows the current Beam object.',4.2)
page('Know what you are editing','Read the context at the top of the inspector before changing a value. Blender selection and Beam’s working context can differ.',[
('table',[['Context','Meaning'],['Selected','The native Blender selection; G and R act on this object.'],['Active','The Beam object whose controls are currently displayed.'],['Pinned','The pin keeps Beam focused while you select another scene object.'],['Receiver','The mesh used by the relevant analysis or Reverse Mapping operation.']]),
('steps',['Select PJ02 in the Outliner and open Projector. Confirm Active refers to PJ02.','Click the pin beside Active. Select a screen or ordinary mesh. Beam retains the pinned projector’s controls.','Unpin when you want the inspector to follow your selection again.']),
('note','Surface Analysis has its own Receiver field. Recalling a study view does not automatically choose a different receiver or rerun analysis. Always check the receiver and source names.')],A['context'],'Pinned context keeps the projector available while another object is selected.',4.2)
page('Install and enable Beam','Start with the supplied Beam 0.13.0 install ZIP. Keep it zipped. A source-code archive and the add-on install archive are different files.',[
('steps',['Open Blender Preferences from the application menu on macOS or Edit → Preferences where available. Choose Add-ons.','Open the small menu at the upper right and choose Install from Disk. Select beam-0.13.0.zip and confirm the installation.','Search for Beam in Add-ons. If its checkbox is off, enable it. Save Preferences if automatic preference saving is disabled.','Close Preferences. If you replaced an older Beam build, save your work and restart Blender before continuing. Check that Beam is still enabled.']),
('note','Use the build supplied with this guide. The documentation includes a small workspace-menu correction; an earlier 0.13.0 ZIP can have the same version number but lack the search entry.'),
('p','Checkpoint: Beam appears enabled in Add-ons. A failed installation or missing module is not solved by opening the training file; enable the add-on first.')],A['install'],'Enabled Beam in a disposable Blender configuration.',6.2)
page('Open the workspace and set the theme','Place the pointer in the 3D View. Press F3 and search for Open Beam Workspace, then select the matching command.',[
('steps',['On a Mac where F3 controls the desktop, use Blender’s Edit → Menu Search instead. You can also use the 3D View’s View → Open Beam Workspace.','In the Beam Inspector expand Project / Workspace. Enter project information such as venue and revision.','Click Apply Beam Theme and confirm the theme change. To return to your previous appearance, click Restore Previous Theme.','Use File → Save As and choose a project folder and .blend filename. Save a working copy of the training project before experimenting.']),
('p','The Beam workspace is a Blender workspace tab. It provides the working layout; the theme changes appearance. Applying the theme is optional. The original layout remains available as another workspace.'),
('note','No search result? Confirm Beam is enabled, restart after replacing an older build, and verify you installed the guide’s build. A sidebar Project / Workspace button is also available.')],A['search'],'Normal Menu Search finds the existing workspace command in the corrected baseline.',6.2)
page('Apply and restore the Beam theme','The theme is optional and affects every Blender workspace. Beam keeps the previous appearance available for restoration.',[
('steps',['Expand Project / Workspace in the inspector. Click Apply Beam Theme.','Read the confirmation and choose Apply Theme. The workspace adopts Beam’s restrained blue/charcoal appearance.','Choose Restore Previous Theme in the same section to restore the captured appearance.','Save the project with File → Save As; a workspace layout and scene belong to the .blend, while application preferences are separate.']),
('note','Use a disposable configuration for training if you do not want to alter your everyday Blender preferences.')],A['theme'],'Project / Workspace contains both theme actions; applying the theme asks for confirmation.',5.6)
page('Create your first screen','Begin in an empty working scene. The example below uses metres and places the centre of a 6 m screen eight metres in front of a projector at the origin.',[
('steps',['Choose Screen → New Screen. In the draft set Width to 6, Height to 3.375 and Aspect to Free. Keep Origin at Bottom Centre. Click Create Screen.','Confirm the new screen is SCR001. In the 3D View use G, Y, 8, Enter; then G, Z, -1.6875, Enter. These offsets put its centre at projector height.','After editing an existing screen’s Width, Height or Aspect, click Update Screen to rebuild it. Moving the object uses Blender’s normal transform tools.','Under Appearance / This Screen choose Test Pattern or Neutral. Use Screen ID, Dimensions and Border to control its labels.']),
('p','Origin matters: Bottom Centre is the bottom-edge origin, not the image centre. Object Centre aiming uses the object’s origin. Use Click Surface when you want to aim at a particular visible point.'),
('note','The screen is receiver geometry. Its dimensions do not force the projected image to match it. Change the projector or use the solver to fit the image.')],A['screen'],'Screen creation settings in the clean-install walkthrough.',3.0)
page('Create your first projector','Choose Projector and click New Projector. Beam assigns PJ identifiers in sequence. Select the projector body or its PJ entry in the Outliner.',[
('steps',['For the first example, keep the new projector at the origin, facing the screen on +Y. Confirm the inspector says TARGET.','Set X to 1920 and Y to 1080 in Projection. Set Throw Ratio to 1.00. Leave H Shift % and V Shift % at zero.','Choose Calibration Grid in Output. Study Colour identifies this projector’s helpers.','Move with G and an axis, or rotate with R and an axis in the 3D View. Press Escape to cancel a transform. Use Aim / Position for deliberate targeting.']),
('p','Brightness is under Advanced as lumens. Nominal illuminance is an estimate based on the nominal image area, not a measured venue brightness result. Surface reflectance, ambient light and optical losses still require on-site judgement.'),
('note','Check Selected before using Blender transforms, and Active before editing Beam controls. A pinned inspector can intentionally show a different object.')],A['projector'],'Projection controls set raster, throw ratio, shift and test output.',3.1)
page('Aim at receiving geometry','The centre ray follows the projector’s optical direction. Beam uses its first eligible scene-geometry hit for the target and nominal calculated values.',[
('steps',['Select the projector, expand Aim / Position and choose Click Surface. Click the intended point on the screen.','Use Object Centre only when the receiver’s object origin is the desired aiming point. 3D Cursor aims at Blender’s cursor.','Use Look Through to inspect the projector’s view. Recheck TARGET and the centre-ray hit after moving either projector or geometry.']),
('p','The target marker identifies the centre hit. A blocker can become that first hit and shorten the nominal throw, even while other rays continue to reach other surfaces. Beam’s actual projection can cover multiple surfaces and respects occlusion.'),
('p','NO TARGET means no eligible centre-ray hit was found. Aim at receiving geometry and check visibility. A displayed fallback helper length is not a measured target distance.'),
('note','Frustum helpers are visual aids. Corner rays stop at their first scene hit or a fallback display distance. Their endpoints do not redefine target calculations or multi-surface projection.')],A['overview'],'Three projectors aimed at a receiving screen; helpers reveal their geometry.',6.8)
page('Read the calculated values','Calculated values describe the nominal centre-target image. Surface Analysis later measures local conditions on a chosen receiver.',[
('table',[['UI label','What it tells you'],['Throw','Lens-to-target distance along the centre ray.'],['Width and Height','Nominal image dimensions at that target.'],['Ratio','Throw distance divided by image width.'],['Density','Nominal horizontal pixels per metre.'],['Pixel Size','Nominal millimetres per pixel; smaller means finer sampling.']]),
('p','Worked example: 1920 × 1080 at an 8 m throw with ratio 1.00 gives an 8 × 4.5 m image, 240 px/m and 4.17 mm/px. A 6 m width at the same throw needs ratio 1.333 and gives 320 px/m, 3.12 mm/px.'),
('p','The m / mm control changes displayed length units. Pixel Size remains a pixel-related measurement. Advanced also exposes DPI and estimated illuminance where applicable.'),
('note','An oblique or curved receiver has different local density and stretch across its surface. Do not treat the nominal panel as proof of uniform image quality.')],A['calc'],'Verified six-metre result after applying the inverse solver.',3.2)
page('Solve for a six metre image','Use this tutorial on the first-projector scene: 1920 × 1080, an 8 m centre throw and a screen facing the projector.',[
('steps',['In Calculated, click the pencil beside Width and enter 6 m as the desired value. This opens a draft solver edit.','Choose Lens / Throw Ratio in Solve By. Keep Target locked so the intended target stays fixed; leave Lens unlocked.','Review the lock controls. Position controls projector movement, Lens controls throw-ratio changes, and Target protects the target.','Click the checkmark to Apply Target. Use the X to cancel the draft.','Inspect the result: Throw 8.00 m, Width 6.00 m, Height 3.38 m and Ratio 1.333.']),
('p','Projector Position changes placement; Lens / Throw Ratio changes the ratio; Both / Auto allows both within the available constraints. Choose deliberately according to what can move in the venue.'),
('note','If constraints conflict, Beam refuses the solution. Check the target, locks, selected lens range and shift. Correct the constraint and try again; do not assume a failed Apply changed the study.')],A['solver'],'Draft solver controls. The training scene uses a different 6.4 m throw.',3.0)
page('Use real optics guidance','Generic / Manual lets you enter values directly. Presets provide known projector and lens data; they do not certify an entire installation.',[
('steps',['Select a projector and expand Optics Guidance. Choose the projector preset and compatible Lens.','Click Use Preset Values to apply the preset raster and lumens. Confirm those values suit your study.','Read the displayed Throw range and status. Set a ratio inside the supported range.','Check the H and V shift information before asking the solver to use shifted optics.']),
('p','The training project’s PJ01 uses Panasonic PT-RZ990 with ET-DLE105 from Beam’s existing verified library. PJ02 and PJ03 are Generic / Manual examples. Their matched 1920 × 1200 raster keeps the blend consistent.'),
('note','Shift limits unavailable means Beam does not know that lens envelope. It deliberately does not invent it. Non-zero unknown shift can prevent a constrained solve; manual editing remains available. Obtain manufacturer data before relying on that configuration.')],A['optics'],'The frozen baseline’s real optics panel and known-limit feedback.',3.0)
page('Control helpers and output','Use annotations to keep a technical view legible. The projected name, body/measurement labels have separate visibility controls.',[
('steps',['Open the projector’s Annotations section. Toggle Body / Measurement Labels independently from Projected Name. The former controls the body ID and measurements together.','Use Show Frustum, Centre Ray and Target Marker when checking aim. Hide them when they obscure receiver content.','Use Show Grid and Projection Output according to the test output you need. Output selects the pattern; visibility controls determine what is drawn.','Use Global · Viewport Display for study-wide presentation settings. Use Technical for inspection and Presentation for a cleaner review.']),
('p','Available outputs include Solid Colour, Calibration Grid, Identifier, Checkerboard, Blender UV Grid, Blender Color Grid and Custom Image. A custom projected image is distinct from Reverse Mapping, which uses the receiver’s existing UV/material texture.'),
('note','A helper being hidden does not remove the projector or change engineering calculations. For exported study views, check This View · Export Annotations as well as the live viewport.')],A['overview'],'Keep the helpers that answer your current question; remove visual clutter for review.',6.8)
page('Build a blend group','A blend group holds editable projectors together. Use it for coordinated placement and overlap planning.',[
('steps',['Select the projectors you want to group in Blender; make the intended anchor the active selection. Choose Blend → From Selection.','Confirm BG01 and its members. The anchor provides the layout reference. Select the group controller to transform the arrangement together.','Use the Members controls to inspect, add, remove or create members. Individual offsets allow local adjustments relative to the group.','Set the requested horizontal or vertical overlap. Compare the actual result in pixels, metres and percent. Expand Pair Details for individual joins.']),
('p','The training project has PJ01, PJ02 and PJ03 in BG01. Its initial horizontal overlap is 240 px, 0.800 m and 12.5%. Reflowing at 15% gives 288 px and 0.960 m for the matched 1920-pixel images.'),
('note','Requested overlap is your intention; actual overlap is the resulting geometry. Check both, especially after member offsets or placement changes.')],A['blend'],'BG01 after the verified overlap change and Reflow Group.',3.0)
page('Create a new automatic layout','New Auto Blend creates new projector objects. It is the starting point for a new layout on a screen.',[
('steps',['Select SCR001. For this example, first resize the screen as described below, then open Screen → New Auto Blend.','For the first-projector example, enlarge the screen to 18 m wide, keep height 3.375 m and click Update Screen.','Choose a horizontal layout with Count 3, Auto Count off, overlap 200 px, Use Active Projector Template on and Lock Throw off. Use the 1.333-ratio practice projector as the template.','Check Valid and coverage before Create Layout. The tested result uses about 8.597 m throw with 0.672 m physical overlap.']),
('p','The existing practice projector remains in the scene; this command creates additional projectors. Hide or remove that practice object when you want only the new layout, and check which projectors are included in analysis and exports.'),
('note','Building a new layout? Use New Auto Blend. Already have BG01? Use Reflow Group. A 6 m screen with this three-projector/template setup was rejected; changing screen size is a deliberate design change, not an automatic fix.')],A['auto'],'The planner previews validity and coverage before creating objects.',3.5)
page('Reflow an existing blend','Reflow Group recalculates member placement in the selected group. It does not create a second independent layout.',[
('steps',['Open the supplied training project and select BG01 through the Outliner or Blend context.','Change requested horizontal overlap from 12.5% to 15%.','Click Reflow Group. Inspect the new positions and the requested/actual readouts.','Confirm three members remain and the result shows 288 px, 0.960 m and 15%. Check the projected joins in the 3D View.']),
('p','Group transforms move the arrangement. Member offsets adjust individual projectors. Reflow changes placement according to the current group settings, so save a working copy before testing a different layout.'),
('p','Overlap previews and feathering help review the plan. They do not replace calibration, warp or blend setup in the real playback system.'),
('note','If coverage is incomplete or constraints are invalid, review count, throw and overlap. A layout that looks plausible still needs valid results and receiver coverage checks.')],A['blend'],'Requested and actual overlap should agree in this matched training example.',3.1)
page('Save and recall study views','Study Views make repeatable review angles. They store framing and export settings, not a new copy of the engineering scene.',[
('steps',['Navigate the 3D View to a useful angle. Choose Views → New Study View, enter a descriptive name and confirm.','Use View beside an entry to recall it. Exit leaves camera view; Previous and Next step through entries.','Set Width and Height for the export image. Open This View · Export Annotations and choose the helpers to include.','Click Export This View and choose the destination. Use Export → Export All Study Views for a set.']),
('p','The training file includes Overview, Front, Blend, Reverse Mapping, Density and Stretch. The names describe useful angles. The Density and Stretch entries also request those analysis overlays when exported.'),
('note','Recalling a view does not reanalyse geometry or change your analysis receiver. Before exporting an analytical view, set the receiver and sources, then Analyse Surface. Technical and Presentation are display choices, not separate engineering studies.')],A['views'],'Named views provide repeatable framing and image export dimensions.',3.1)
page('Analyse a receiver surface','Surface Analysis samples points across a selected mesh and checks which projectors can actually reach each point.',[
('steps',['Expand Surface Analysis. Set Receiver explicitly to SCR001.','Choose Projectors: Active Projector, Active Blend Group or All Projectors. For the training blend use Active Blend Group and confirm BG01 is the active group.','Choose Samples: Quick · 512, Standard · 2048 or Detailed · 8192. Start with Standard.','Click Analyse Surface. Confirm the receiver and source names listed below the button.','Enable Show Heatmap and select Density or Stretch. Read the statistics together with the spatial overlay.']),
('p','More samples give finer inspection but take more work. This is sampled coverage, not an exhaustive proof of every tiny gap. Thin obstacles and areas between samples can be missed.'),
('note','After moving a projector, receiver or obstruction, the result can become stale. Scene changed · Analyse again means the old result must be refreshed before you rely on it or export it.')],A['density'],'A refreshed SCR001 result with all three BG01 members as sources.',3.0)
page('Interpret density','Density describes how much projected raster detail reaches the receiver. Higher px/m means more pixels across the same physical distance.',[
('table',[['Readout','How to use it'],['Min density','Inspect the weakest sampled area.'],['Avg density','Summarises visible sampled coverage; does not describe every point.'],['Max density','Shows the densest sampled area.'],['Visible coverage','Receiver area reached by at least one visible source.'],['Overlap','Receiver area with more than one visible source.']]),
('p','The default density colour scale is local to the current analysis: dark blue is its low end and pale teal its high end. Two analyses can show similar colours while having very different numerical densities. Compare the values, not colour alone.'),
('p','On the front-facing training screen the visible samples are approximately 300 px/m. The curved receiver varies strongly with angle. Select it and analyse again to compare.'),
('note','Beam selects the best visible projector by conservative worst-direction density at each sample. Overlap is not simply summed as extra image resolution.')],A['density-overlay'],'Actual Density export after the blocker move. The legend reports the numerical range.',5.3)
page('Inspect stretch on curved geometry','A pixel projected straight onto a surface is relatively even. At a steep angle it spreads unevenly. Stretch highlights that directional distortion.',[
('steps',['Select TRAINING_Curved_Receiver as the Surface Analysis receiver. Keep the appropriate blend sources.','Click Analyse Surface and choose Stretch in Heatmap.','Recall the Stretch study view to examine the curved surface at a useful scale.','Look for changes from teal near 1:1 toward amber near 3:1. Compare the angle of the surface with the projector direction.']),
('p','Stretch is anisotropy: how unevenly the pixel footprint expands in different directions. It is not the same as low density. A surface can have many pixels per metre and still show directional stretching.'),
('p','For the validated training sequence, the curved receiver showed 100% sampled coverage and density from about 214 to 384 px/m. Those numbers apply to that geometry and layout, not every curved screen.'),
('note','The exported view uses the latest analysis receiver. Its name “Stretch” does not automatically choose the curved mesh.')],A['stretch'],'Actual exported Stretch overlay on the training receiver.',6.8)
page('Find occlusion with the blocker','Use this controlled comparison after reflowing the training group to 15%. Keep receiver SCR001, Active Blend Group and Standard · 2048 samples.',[
('steps',['With TRAINING_Occluder at its starting X = -9 m, Analyse Surface. The tested visible coverage is 93.1%; the curved receiver already obstructs part of the screen.','Select TRAINING_Occluder in the Outliner. In the 3D View use G, X, 9, Enter to move it to world X = 0 m.','Observe Scene changed · Analyse again. Click Analyse Surface again.','Compare: visible coverage falls to 78.6%, and occluded samples rise from 6.9% to 21.4%.']),
('p','Nominal coverage ignores blocking geometry. Visible coverage tests whether a ray reaches the receiver first. Here nominal screen coverage remains 100%, while the blocker removes visible coverage.'),
('note','A blocker in an overlap zone may not reduce total visible coverage: another projector can still reach that area. Move this blocker to the centre beam for the clear comparison. Results are sampled and may vary with changed settings.')],A['comparison'],'Same receiver and sample settings; the blocker moves into the centre beam.',5.2)
page('Reveal mapped content with Reverse Mapping','Reverse Mapping uses the receiver’s existing UV/material texture and reveals that texture according to actual projector coverage. It helps answer which parts of mapped content are reached.',[
('steps',['Recall the Reverse Mapping study view. Select SCR001 in the Outliner. Use Material Preview shading in the 3D View so the texture is visible.','Expand Reverse Mapping and confirm Receiver is SCR001. The supplied screen already has valid UVs and a packed colour-grid texture.','Set Coverage From to Blend Group, then choose BG01. Other choices include Active Projector, Selected Projector and All Active Projectors.','Click Enable on Receiver. Inspect the receiver content and the shaded or missing regions around occlusion.','Click Restore Materials when finished to return the receiver’s original material state.']),
('p','The projected Calibration Grid belongs to the projector. The coloured letter chart in this example belongs to the receiver. Changing Reverse Mapping coverage reveals that same receiver content; it does not move the texture into projector space.'),
('note','Hide unnecessary projection helpers while judging mapped content. Keep the receiver and source names visible so a clean view is not mistaken for a different study.')],A['reverse-controls'],'Receiver, Coverage From, Mode and BG01 in the actual Reverse Mapping panel.',4.1)
page('Choose a Reverse Mapping mode','Choose the display that best answers your review question. The coverage sources stay the same while the treatment of covered and uncovered areas changes.',[
('table',[['Exact mode','Practical use'],['Covered Only','Show covered texture against neutral uncovered regions.'],['Uncovered Dimmed','Keep missed content visible at Outside Brightness for comparison.'],['Uncovered Hidden','Make uncovered parts transparent to reveal missed geometry.'],['Overlap Diagnostic','Mark overlap in amber while reviewing content and coverage.']]),
('p','Try each mode with the blocker at X = 0 m. The missed centre region is easy to compare. In Overlap Diagnostic the vertical overlap bands identify where adjacent projectors both contribute.'),
('note','Transparency in Uncovered Hidden is a review treatment. It does not physically remove the receiver or prove that transparent real-world materials will behave the same way.')],A['mode-grid'],'Same receiver and sources: compare the central blocked region and the overlap bands.',6.8)
page('Use Reverse Mapping within its limits','Reverse Mapping is an interactive GPU viewport tool. Use it to review coverage, mapped content, missed geometry, overlap and alignment.',[
('steps',['Use a normal GPU-backed Blender viewport and Material Preview. Allow shaders to finish compiling after changing mode.','Supply a supported node material with a usable image texture and valid UVs. The training screen provides a known working starting point.','Choose a source set containing valid projectors. The viewport supports up to 32 coverage sources. Reduce a larger set to the group being reviewed.','After an error, read the receiver and coverage status before enabling again. Restore Materials returns the original material assignment.']),
('p','Coverage depth is a viewport approximation with material and UV assumptions. Transparent and unusual shader setups should not be assumed equivalent to opaque receiving surfaces. This is not an offline final-rendering system.'),
('p','Use Surface Analysis for sampled quantitative statistics, Reverse Mapping for content-aware coverage review, and actual installation tests for final confidence.'),
('note','These workflows were verified on macOS with Blender 5.2.2 LTS. This package does not claim Windows, Linux or a different GPU configuration has been tested.')],A['hidden'],'Uncovered Hidden reveals the geometry behind a missed portion of the receiver.',6.8)
page('Choose the correct export','Mapping Matter CSV and Designer CSV are separate formats. Pick the button for the receiving workflow; renaming one CSV does not convert it into the other.',[
('table',[['Export button','Destination and purpose'],['Mapping Matter CSV','Mapping Matter-format projector data.'],['Designer CSV','Designer projector-study CSV workflow.'],['Beam JSON','Beam technical study data.'],['Study Package · JSON + HTML + Views','Technical handoff bundle for PATCH integration and review.']]),
('steps',['Open Export and choose the exact format. Set a distinct destination filename and confirm.','Keep the generated .metadata.json sidecar with each CSV. It records export context and conventions.','Check the file exists and contains the expected projector rows. For the training project, check PJ01, PJ02 and PJ03.']),
('p','Screen Geometry offers OBJ and FBX separately from projector data. Check the screen, production surface choice and coordinate mode before exporting geometry.'),
('note','The two CSV outputs were written and checked in this walkthrough. Import into external Mapping Matter or Designer applications was not part of this validation.')],A['export'],'Separate buttons prevent accidental substitution of the two CSV formats.',3.6)
page('Export analytical images and hand off to PATCH','Use a named study view for repeatable framing and a Study Package for a complete technical handoff.',[
('steps',['For a Density image, set Surface Analysis Receiver to SCR001, choose the sources and Analyse Surface.','Recall Density. Under This View · Export Annotations confirm Analysis Overlay is Density. Choose labels and other study overlays deliberately.','Click Export This View, choose a folder and export. Keep the image and its .analysis.json sidecar together.','Open Export → Study Package · JSON + HTML + Views. Enter a new folder name such as Training_Study_Package and leave Render Study Views on. Confirm the export.','Check beam-study.json, index.html and the views folder. Open index.html for a review of the handoff.']),
('p','Intended handoff: Beam → Study Package → PATCH → Documents / client-facing reporting. Beam owns technical study data. PATCH owns document composition.'),
('note','PATCH import integration is still being tested. This guide validates Beam’s package output, not a complete PATCH GUI import or final client document workflow. Beam is not a report-layout application.')],A['export-density'],'Density is an explicit export overlay choice; a view name alone is insufficient.',3.4)
page('Troubleshoot the current study','Start by checking the active object, receiver and source names. Save before changing a working layout.',[
('table',[['Message or symptom','Meaning and action'],['NO TARGET','No eligible centre hit. Aim at receiving geometry; check visibility and target marker.'],['Scene changed · Analyse again','Cached analysis no longer matches the scene. Re-run Analyse Surface.'],['Surface Analysis is absent or stale. Analyse again before exporting.','Refresh the intended receiver/source set before exporting an analysis view.'],['Selected lens supports … throw. Current: …','The ratio is beyond known limits. Select a compatible lens or a valid ratio.'],['Shift limits unavailable','Beam has no verified shift envelope. Do not assume a constrained shifted solve is supported.'],['No visible coverage source','Choose a valid projector/group source set; check its members.'],['Reverse Mapping needs an interactive GPU viewport','Use an interactive GPU viewport. A background process cannot provide the same viewport path.'],['Unsupported receiver material','Use a supported UV/image node material; compare the training screen.'],['Export fails','Read the error, choose a writable destination, refresh analysis if needed and retry. Verify actual output files.'],['Wrong controls or object changes','Compare Selected and Active. Unpin if you want the inspector to follow selection.'],['Beam rail missing','Pointer over 3D View → T. Alternatively use the five inspector module buttons.']]),
('note','If Open Beam Workspace is absent from search, verify the supplied corrected build and enabled add-on. Record the build, Blender version and exact message when reporting an issue.')])
page('Where is it','Use this page as a working reference. Paths start in the Beam Inspector unless stated otherwise.',[
('table',[['Task','Location'],['Add projector','Projector → New Projector'],['Add screen','Screen → New Screen → Create Screen'],['Resize screen','Screen → Width / Height → Update Screen'],['Make a group from projectors','Blend → From Selection'],['Create a new layout','Screen or Blend → New Auto Blend'],['Adjust an existing layout','Blend → Reflow Group'],['Change lens','Projector → Optics Guidance'],['Solve image size','Projector → Calculated pencil → Solver'],['Aim at a visible point','Projector → Aim / Position → Click Surface'],['Hide projected name','Projector → Annotations → Projected Name'],['Save or recall framing','Views → New Study View / View'],['Density analysis','Surface Analysis → Analyse Surface → Heatmap: Density'],['Stretch analysis','Surface Analysis → Heatmap: Stretch'],['Reveal receiver texture','Reverse Mapping → Enable on Receiver'],['Mapping Matter data','Export → Mapping Matter CSV'],['Designer data','Export → Designer CSV'],['PATCH handoff bundle','Export → Study Package · JSON + HTML + Views'],['Theme and workspace','Project / Workspace'],['Save the Blender project','Blender File → Save As / Save']]),
('p','Next: open Beam_Training_Project.blend and follow the separate training companion. Work on a copy, and return to the original file to reset all exercises.')])

from content_013 import extend,training_additions
extend(pages,page,A,out)

# DOCX layout and reusable rendering helpers.
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
for n in ['Normal','Body Text']:
 s=doc.styles[n];s.font.name='Arial';s.font.size=Pt(10.5);s.font.color.rgb=RGBColor.from_string(INK);s.paragraph_format.space_after=Pt(7);s.paragraph_format.line_spacing=1.08
for n,size in [('Title',29),('Heading 1',23),('Heading 2',13)]:
 s=doc.styles[n];s.font.name='Arial';s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=RGBColor(0,0,0);s.paragraph_format.space_after=Pt(10)
for style in doc.styles:
 for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
foot=sec.footer.paragraphs[0];foot.text='Beam 0.13.0 — Internal Production Beta   •   Twisted Pixel';foot.runs[0].font.size=Pt(8)
foot.add_run('\nBlender 5.2.2 LTS • macOS 26.7                                      ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
for r in foot.runs:r.font.size=Pt(8)
doc.core_properties.title='Beam User Guide';doc.core_properties.author='Twisted Pixel';doc.core_properties.subject='Beam 0.13.0 onboarding and operation'
def para(txt,style=None):
 p=doc.add_paragraph(txt,style);p.paragraph_format.widow_control=True;return p
def table(rows):
 t=doc.add_table(rows=0,cols=len(rows[0]));t.autofit=False;t.columns[0].width=Inches(2.0);t.columns[1].width=Inches(4.7)
 for i,row in enumerate(rows):
  cs=t.add_row().cells
  for j,txt in enumerate(row):
   cs[j].text=txt;tcPr=cs[j]._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),INK if i==0 else ('F0F3F5' if i%2 else 'FFFFFF'));tcPr.append(shade)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tcPr.append(borders)
   for p in cs[j].paragraphs:
    p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
    for r in p.runs:r.font.size=Pt(9.5);r.font.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else INK)
  trPr=t.rows[-1]._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
 para('')
for idx,p in enumerate(pages):
 if idx:doc.add_page_break()
 para(p['title'],'Title' if idx==0 else 'Heading 1');para(p['intro'])
 if p['img']:
  ip=doc.add_paragraph();ip.alignment=1;ip.add_run().add_picture(p['img'],width=Inches(p['width']))
  cp=para(p['caption']);cp.alignment=1
  for r in cp.runs:r.italic=True;r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('52616A')
 for typ,val in p['blocks']:
  if typ=='table':table(val)
  elif typ=='steps':
   for i,step in enumerate(val,1):para(f'{i}. {step}')
  elif typ=='note':
   np=para('CHECK  '+val);np.paragraph_format.space_before=Pt(4)
   for r in np.runs:r.font.color.rgb=RGBColor.from_string(AMBER);r.font.size=Pt(10)
  else:para(val)
doc.save(out/'Beam_User_Guide.docx')
(out/'manual-content.json').write_text(json.dumps(pages,indent=2))

# Companion PDFs use a controlled page layout and the same verified source images.
styles={
 'title':ParagraphStyle('title',fontName='Helvetica-Bold',fontSize=24,leading=28,spaceAfter=14,textColor=colors.HexColor('#'+INK)),
 'h':ParagraphStyle('h',fontName='Helvetica-Bold',fontSize=14,leading=18,spaceBefore=12,spaceAfter=6,textColor=colors.HexColor('#'+INK)),
 'body':ParagraphStyle('body',fontName='Helvetica',fontSize=10.5,leading=14,spaceAfter=7,textColor=colors.HexColor('#'+INK)),
 'cap':ParagraphStyle('cap',fontName='Helvetica-Oblique',fontSize=9,leading=12,spaceAfter=10,textColor=colors.HexColor('#52616A'))}
def pp(t,style='body'):return Paragraph(t.replace('&','&amp;'),styles[style])
def ri(key,w=480,maxh=230):
 im=Image.open(A[key]);iw,ih=im.size;scale=min(w/iw,maxh/ih);return RImage(A[key],iw*scale,ih*scale)
def footer(c,d):
 c.setFillColor(colors.HexColor('#52616A'));c.setFont('Helvetica',8);c.drawString(44,30,'Beam 0.13.0 — Internal Production Beta  |  Blender 5.2.2 LTS  |  macOS 26.7');c.drawRightString(550,30,str(d.page))
def pdf(name,story):SimpleDocTemplate(str(out/name),pagesize=(595.28,841.89),rightMargin=44,leftMargin=44,topMargin=40,bottomMargin=48).build(story,onFirstPage=footer,onLaterPages=footer)
q=[pp('Beam Quick Start','title'),pp('Twisted Pixel • A first projection study in two pages'),ri('modules',400,115),pp('1  Install and open Beam','h'),pp('Install beam-0.13.0.zip from Preferences → Add-ons → menu → Install from Disk. Enable Beam and save preferences if needed. Restart after replacing an older build. In the 3D View press F3, search Open Beam Workspace and run it. On Mac use Edit → Menu Search if F3 is intercepted. View → Open Beam Workspace is also available.'),pp('2  Add and position a screen','h'),pp('Start in an empty scene: in a new General file, select the default objects in the 3D View with A and delete them with X, then confirm. Choose Screen → New Screen. Set Width 6, Height 3.375, Aspect Free and Origin Bottom Centre; click Create Screen. With SCR001 selected, use G Y 8 Enter, then G Z -1.6875 Enter. The screen centre is now eight metres along +Y at projector height.'),pp('3  Add a projector and check its target','h'),pp('Choose Projector → New Projector. Keep it at the origin facing +Y. Use resolution 1920 × 1080, Throw Ratio 1.00 and zero shift. Choose Calibration Grid. Confirm TARGET; Calculated should read Throw 8 m and Width 8 m. If needed use Aim / Position → Click Surface to aim at the screen centre.'),pp('4  Make the image six metres wide','h'),pp('Click the Width pencil, enter 6 m, choose Solve By: Lens / Throw Ratio, keep Target locked and Lens unlocked, then click the checkmark. Confirm Width 6 m and Ratio 1.333.'),pp('Optional appearance: Project / Workspace → Apply Beam Theme. Restore Previous Theme reverses it. File → Save As saves your study.','cap'),PageBreak(),pp('Build and hand off a layout','title'),ri('auto',360,220),pp('5  Create a new three-projector layout','h'),pp('For this example, select SCR001, change Width to 18 m and click Update Screen. Choose New Auto Blend: Horizontal, Count 3, Auto Count off, overlap 200 px, Use Active Projector Template on, Lock Throw off. Keep the 1.333 practice projector active as the template. Confirm Valid, then Create Layout. The tested throw is about 8.597 m.'),pp('New Auto Blend adds projectors; the original practice projector remains. Hide or remove it when you want only the layout. For an existing BG01 use Blend → Reflow Group instead of creating another layout.','cap'),pp('6  Save a review view','h'),pp('Navigate to show the projectors and screen. Choose Views → New Study View and name it Overview. Use View to recall it. Set export dimensions and inspect This View · Export Annotations. Save the .blend.'),pp('7  Export and verify','h'),pp('Choose Export → Study Package · JSON + HTML + Views. Enter a new folder name, leave Render Study Views on and export. Check beam-study.json, index.html and the views folder. Use Mapping Matter CSV only for that format; use Designer CSV for the Designer workflow. Keep metadata sidecars.'),pp('Next: open the supplied training project and follow its companion for all five Surface Analysis modes, occlusion, Observer Acuity and Reverse Mapping. PATCH document composition is a separate integration workflow still in testing.','cap')]
pdf('Beam_Quick_Start.pdf',q)
training=[
('Start with a working copy','Open Beam_Training_Project.blend with the guide’s Beam build enabled. File → Save As to a new name. Use metres. The original file resets the exercises; analysis caches are recalculated after opening.', 'overview',[
('1  Find every module','Locate Projector, Screen, Blend, Views and Export in the inspector and left rail. Press T over the 3D View if the rail is hidden. Expand BEAM collections in the Outliner.'),
('2  Select PJ02','Select PJ02 from BEAM → Projectors. Confirm Active: PJ02 in Projector. It is a Generic / Manual example; PJ01 uses the existing Panasonic PT-RZ990 / ET-DLE105 preset.'),
('3  Move PJ02','With pointer in the 3D View press G, X, 1, Enter. Observe the moved image. Undo with Command-Z on Mac (Ctrl-Z on other Blender setups) before continuing. Check that PJ02 returned.')]),
('Change the image and solve','The starting raster is 1920 × 1200. PJ02 starts at approximately 6.4 m throw with ratio 1.00 and image width 6.4 m.', 'training-solve',[
('4  Change throw ratio','In Projection set Throw Ratio to 1.20. Confirm Width about 5.33 m, Height about 3.33 m and Density about 360 px/m.'),
('5  Use the inverse solver','Click the Width pencil and enter 6 m. Solve By: Lens / Throw Ratio. Keep Target locked and Lens unlocked. Apply with the checkmark. Confirm Width 6 m, Height 3.75 m, Throw 6.40 m and Ratio about 1.067. Restore Throw Ratio to 1.00 after this exercise.'),
('6  Adjust BG01 overlap','Select BG01 or open its Blend context. Change requested horizontal overlap from 12.5% to 15%. This is the desired join; inspect the result after reflow.')]),
('Reflow and measure coverage','Keep the changed 15% overlap for the following comparison. Use the same receiver, source set and sample count for both measurements.', 'blend',[
('7  Reflow BG01','Click Reflow Group. Confirm PJ01, PJ02 and PJ03 remain members. Requested and actual horizontal overlap should read 288 px, 0.960 m and 15%.'),
('8  Run Density analysis','Expand Surface Analysis. Set Receiver SCR001, Projectors Active Blend Group and Samples Standard · 2048. Click Analyse Surface. Set Heatmap to Density. Confirm sources PJ01, PJ02, PJ03; visible coverage is about 93.1% with the blocker at X = -9 m. Visible samples are about 300 px/m.'),
('9  Move the blocker','Select TRAINING_Occluder in the Outliner. In the 3D View press G, X, 9, Enter. It moves from X = -9 to X = 0 m. Leave it there for the next exercises.')]),
('Compare occlusion and stretch','The geometry has changed. The old analysis is stale and must be refreshed.', 'after',[
('10  Reanalyse and compare','With SCR001 and the same group/settings, click Analyse Surface again. Visible coverage falls to about 78.6%; occluded samples rise to 21.4%. Nominal coverage remains 100%. A blocker at X = -2 lies in an overlap region and can leave total coverage unchanged; use X = 0 for this demonstration.'),
('11  Inspect Stretch','Choose Stretch in Heatmap. The flat screen is close to 1:1. Then set Receiver to TRAINING_Curved_Receiver and Analyse Surface. Recall the Stretch view for a close angle. Inspect the changing surface angle and overlay; the tested curved result has 100% sampled coverage and roughly 214–384 px/m density. Turn Show Heatmap off before reviewing the receiver texture.')]),
('Review mapped content','The colour chart is packed into the .blend and the receiver has valid UVs. Use the receiver’s material in Material Preview.', 'overlap',[
('12  Enable Reverse Mapping','Recall the Reverse Mapping study view. Select SCR001 in the Outliner. Expand Reverse Mapping, confirm Receiver SCR001, choose Coverage From: Blend Group and BG01. Click Enable on Receiver. Use the Material Preview sphere at the upper right of the 3D View.'),
('13  Try all four modes','Compare Covered Only, Uncovered Dimmed, Uncovered Hidden and Overlap Diagnostic. With the blocker at X = 0, the central missed region is visible. Dimmed retains content at Outside Brightness; Hidden reveals geometry behind it; Overlap Diagnostic marks joins amber. Finish with Restore Materials.'),
('14  Recall Study Views','Use View beside Overview, Front, Blend, Reverse Mapping and Density; scroll the list to Stretch. Use Exit to leave camera view. View recall does not change the analysis receiver or rebuild its result.')]),
('Export the training study','Use distinct filenames and keep sidecars with their exports. These steps complete the Beam side of the handoff.', 'export',[
('15  Export a Density view','Set Surface Analysis Receiver back to SCR001, Active Blend Group, Standard · 2048 and Analyse Surface. Recall Density. Confirm This View · Export Annotations → Analysis Overlay: Density. Export This View. Check the PNG and .analysis.json sidecar; the latter records receiver, sources and statistics.'),
('16  Export Mapping Matter CSV','Open Export → Mapping Matter CSV. Save Beam_projectors.csv. Check three PJ rows and keep Beam_projectors.metadata.json.'),
('17  Export Designer CSV','Use the separate Designer CSV button. Save Beam_Designer.csv and keep its metadata sidecar. Its columns differ from Mapping Matter; do not substitute one for the other.'),
('18  Export a Study Package','Choose Study Package · JSON + HTML + Views. Enter a new folder name, leave Render Study Views on and export. Check beam-study.json, index.html and nine view images. The JSON includes the technical study; PATCH document composition is a separate integration still in testing.')])]
training.extend(training_additions())
story=[]
for i,(title,intro,key,exs) in enumerate(training):
 if i:story.append(PageBreak())
 story +=[pp('Beam Training Project Guide' if i==0 else title,'title')]
 if i==0:story.append(pp(title,'h'))
 story +=[pp(intro),ri(key,470,190)]
 for h,t in exs:story +=[pp(h,'h'),pp(t)]
 if i==5:story.append(pp('Save your exercise copy. Reopen the original training file to reset projector placement, overlap and blocker position.','cap'))
pdf('Beam_Training_Project_Guide.pdf',story)
print('Created DOCX and two companion PDFs; manual topic pages:',len(pages))
